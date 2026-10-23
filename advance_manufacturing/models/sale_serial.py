import base64
import io

from odoo import api, fields, models, _
from odoo.exceptions import UserError


class AdvmfgSoSerial(models.Model):
    """A serial number promised to a customer's sale order (before or while the panel is built).

    Think of it as a name tag reserved for one sale-order line. If the order is cancelled the tag
    is 'released' and handed to the next order for the same product instead of wasting the number.
    """
    _name = 'advmfg.so.serial'
    _description = 'Sale Order Serial Number'
    _order = 'name'

    name = fields.Char(string='Serial Number', required=True, copy=False)
    sale_order_id = fields.Many2one('sale.order', ondelete='set null', index=True)
    line_id = fields.Many2one('sale.order.line', ondelete='set null')
    product_id = fields.Many2one('product.product', required=True)
    company_id = fields.Many2one('res.company', required=True, default=lambda self: self.env.company)
    production_id = fields.Many2one('mrp.production', string='Manufacturing Order', ondelete='set null')
    state = fields.Selection(
        [('assigned', 'Assigned'), ('delivered', 'Delivered'), ('released', 'Released')],
        default='assigned', required=True)

    _sql_constraints = [
        ('name_company_uniq', 'unique(name, company_id)', 'Serial Number already exists!'),
    ]


class SaleOrder(models.Model):
    """Serial numbers on the sale order: generate, reuse and release."""
    _inherit = 'sale.order'

    advmfg_so_serial_ids = fields.One2many('advmfg.so.serial', 'sale_order_id', string='Serial Numbers')
    advmfg_so_serial_count = fields.Integer(compute='_compute_advmfg_so_serial_count')

    @api.depends('advmfg_so_serial_ids')
    def _compute_advmfg_so_serial_count(self):
        """How many serial numbers this order holds."""
        for order in self:
            order.advmfg_so_serial_count = len(order.advmfg_so_serial_ids)

    def action_advmfg_generate_serials(self):
        """Generate Serials: give each line one serial per unit. Reuse released ones first, then make new."""
        for order in self:
            if order.state == 'cancel':
                raise UserError(_('You cannot generate serial numbers for a cancelled order.'))
            config = self.env['advmfg.serial.prefix'].search([('company_id', '=', order.company_id.id)], limit=1)
            prefix = config.prefix or order.company_id.advmfg_serial_prefix
            if not prefix:
                raise UserError(_('Serial Number Prefix not configured, please configure it first.'))
            Serial = self.env['advmfg.so.serial']
            seq = self.env['ir.sequence']
            for line in order.order_line.filtered(
                    lambda l: l.product_id and l.product_id.type != 'service' and not l.display_type):
                have = Serial.search_count([('line_id', '=', line.id), ('state', '!=', 'released')])
                missing = int(line.product_uom_qty) - have
                if missing <= 0:
                    continue
                # Step 1: reuse numbers freed by cancelled orders for the same product.
                spare = Serial.search([('product_id', '=', line.product_id.id), ('state', '=', 'released'),
                                       ('company_id', '=', order.company_id.id)], limit=missing)
                spare.write({'sale_order_id': order.id, 'line_id': line.id, 'state': 'assigned'})
                # Step 2: whatever is still missing gets brand-new numbers.
                for _i in range(missing - len(spare)):
                    Serial.create({
                        'name': '%s%s' % (prefix, seq.next_by_code('advmfg.serial')),
                        'sale_order_id': order.id, 'line_id': line.id,
                        'product_id': line.product_id.id, 'company_id': order.company_id.id,
                    })
        return True

    def action_view_advmfg_so_serials(self):
        """Smart button: open this order's serial numbers."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': _('Serial Numbers'),
            'res_model': 'advmfg.so.serial', 'view_mode': 'list,form',
            'domain': [('sale_order_id', '=', self.id)],
        }

    def action_cancel(self):
        """After cancelling, free this order's serial numbers so another order can reuse them."""
        res = super().action_cancel()
        for order in self.filtered(lambda o: o.state == 'cancel'):
            order.advmfg_so_serial_ids.filtered(lambda s: s.state == 'assigned').write(
                {'sale_order_id': False, 'line_id': False, 'production_id': False, 'state': 'released'})
        return res


class MrpProduction(models.Model):
    """Manufacturing order extras: note from the sale order, link to a sale order, serial import/export."""
    _inherit = 'mrp.production'

    advmfg_sale_note = fields.Text(string='Additional Notes from SO', compute='_compute_advmfg_sale_note')

    @api.depends('origin')
    def _compute_advmfg_sale_note(self):
        """Show the sale order's extra notes on the MO (matched by the order number in 'Source')."""
        for mo in self:
            order = self.env['sale.order'].search([('name', '=', mo.origin)], limit=1) if mo.origin else False
            mo.advmfg_sale_note = order.note_custom if order else False

    def action_advmfg_link_sale_order(self):
        """Open the 'Link to Sale Order' pop-up."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': _('Link Manufacturing Order to Sale Order'),
            'res_model': 'advmfg.link.sale.wizard', 'view_mode': 'form', 'target': 'new',
            'context': {'default_production_id': self.id},
        }

    def action_advmfg_import_serials(self):
        """Open the 'Import Serial Numbers' pop-up."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': _('Import Serial Numbers'),
            'res_model': 'advmfg.import.serial.wizard', 'view_mode': 'form', 'target': 'new',
            'context': {'default_production_id': self.id},
        }

    def action_advmfg_export_serials(self):
        """Download the serial pool as an Excel sheet."""
        import xlsxwriter
        self.ensure_one()
        buf = io.BytesIO()
        book = xlsxwriter.Workbook(buf, {'in_memory': True})
        sheet = book.add_worksheet('Serial Pool')
        bold = book.add_format({'bold': True})
        for col, head in enumerate(['Serial Number', 'Product', 'Manufacturing Order', 'Status']):
            sheet.write(0, col, head, bold)
        for row, serial in enumerate(self.advmfg_serial_ids, start=1):
            sheet.write_row(row, 0, [serial.name, self.product_id.display_name, self.name, serial.state])
        sheet.set_column(0, 3, 26)
        book.close()
        att = self.env['ir.attachment'].create({
            'name': 'Serial Pool - %s.xlsx' % self.name, 'type': 'binary',
            'datas': base64.b64encode(buf.getvalue()), 'res_model': 'mrp.production', 'res_id': self.id,
        })
        return {'type': 'ir.actions.act_url', 'url': '/web/content/%s?download=true' % att.id, 'target': 'self'}


class AdvmfgLinkSaleWizard(models.TransientModel):
    """Attach an existing manufacturing order to a sale order so the delivery waits for it."""
    _name = 'advmfg.link.sale.wizard'
    _description = 'Link Manufacturing Order to Sale Order'

    production_id = fields.Many2one('mrp.production', required=True)
    sale_order_id = fields.Many2one('sale.order', string='Sale Order', required=True,
                                    domain="[('state', 'in', ('sale', 'done'))]")

    def action_link(self):
        """Link: set the MO source, tie the delivery move to the MO's finished move, set the sale on the group."""
        self.ensure_one()
        mo, so = self.production_id, self.sale_order_id
        if mo.state == 'done':
            raise UserError(_('This Manufacturing Order is already Done - its production is final and can no '
                              'longer be relinked to a different Sale Order.'))
        mo.origin = so.name
        deliveries = so.picking_ids.move_ids.filtered(
            lambda m: m.product_id == mo.product_id and m.state not in ('done', 'cancel') and not m.move_orig_ids)
        finished = mo.move_finished_ids.filtered(lambda m: m.product_id == mo.product_id and m.state not in ('done', 'cancel'))
        if not deliveries:
            raise UserError(_("No open delivery move of '%s' was found on Sale Order %s to link.",
                              mo.product_id.display_name, so.name))
        if not finished:
            raise UserError(_("This Manufacturing Order has no outstanding finished-product move for '%s' to link.",
                              mo.product_id.display_name))
        deliveries.write({'move_orig_ids': [(4, finished[0].id)]})
        finished[0].write({'move_dest_ids': [(4, m.id) for m in deliveries]})
        if mo.procurement_group_id:
            mo.procurement_group_id.sale_id = so.id
        mo.message_post(body=_('Linked to Sale Order %s: source set and %s delivery move(s) now wait on this order.',
                               so.name, len(deliveries)))
        return {'type': 'ir.actions.act_window_close'}


class AdvmfgImportSerialWizard(models.TransientModel):
    """Add serial numbers to an order's pool from an Excel file (column A = serial number)."""
    _name = 'advmfg.import.serial.wizard'
    _description = 'Import Serial Numbers'

    production_id = fields.Many2one('mrp.production', required=True)
    excel_file = fields.Binary(string='Upload Excel File')
    excel_filename = fields.Char()

    def action_download_format(self):
        """Download a blank sheet showing the expected column."""
        import xlsxwriter
        buf = io.BytesIO()
        book = xlsxwriter.Workbook(buf, {'in_memory': True})
        sheet = book.add_worksheet('Serial Import')
        sheet.write_row(0, 0, ['Serial Number', 'Product', 'Manufacturing Order'])
        sheet.set_column(0, 2, 26)
        book.close()
        att = self.env['ir.attachment'].create({
            'name': 'Serial Import Format.xlsx', 'type': 'binary', 'datas': base64.b64encode(buf.getvalue())})
        return {'type': 'ir.actions.act_url', 'url': '/web/content/%s?download=true' % att.id, 'target': 'self'}

    def action_import(self):
        """Read the file and add each new serial number to the pool; duplicates are skipped and counted."""
        import openpyxl
        self.ensure_one()
        if not self.excel_file:
            raise UserError(_('Please upload an Excel file before applying.'))
        try:
            sheet = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.excel_file)), data_only=True).active
        except Exception as e:  # any unreadable file
            raise UserError(_("Could not read the uploaded file. Please make sure it's a valid Excel file. (%s)", e))
        mo = self.production_id
        Serial = self.env['advmfg.serial']
        done = skipped = dup = 0
        seen = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name = str(row[0]).strip() if row and row[0] is not None else ''
            if not name:
                skipped += 1
                continue
            if name in seen or Serial.search_count([('name', '=', name), ('company_id', '=', mo.company_id.id)]):
                dup += 1
                continue
            seen.add(name)
            Serial.create({'name': name, 'production_id': mo.id})
            done += 1
        return {
            'type': 'ir.actions.client', 'tag': 'display_notification',
            'params': {'title': _('Import Complete'), 'sticky': False, 'type': 'success',
                       'message': _('%s serial number(s) imported successfully. %s duplicate(s) and %s blank row(s) skipped.',
                                    done, dup, skipped),
                       'next': {'type': 'ir.actions.act_window_close'}},
        }
